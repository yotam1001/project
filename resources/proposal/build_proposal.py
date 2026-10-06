from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
template = ROOT.parent / 'originals' / 'Template_for_the_high_school_project_proposal (1).docx'
doc = Document(template)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'DejaVu Sans', Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

def para(text='', bold=False, size=11, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    prop = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); bidi.set(qn('w:val'), '1'); prop.append(bidi)
    r = p.add_run(text); r.bold = bold; r.font.name = 'DejaVu Sans'; r.font.size = Pt(size)
    rp = r._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl'); rtl.set(qn('w:val'), '1'); rp.append(rtl)
    cs_size = OxmlElement('w:szCs'); cs_size.set(qn('w:val'), str(size * 2)); rp.append(cs_size)
    if bold:
        cs_bold = OxmlElement('w:bCs'); rp.append(cs_bold)
    fonts = rp.find(qn('w:rFonts'))
    if fonts is not None: fonts.set(qn('w:cs'), 'DejaVu Sans')
    return p

def heading(text): para(text, True, 16)
def page(): doc.add_page_break()
def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    direction = OxmlElement('w:bidiVisual'); t._tbl.tblPr.append(direction)
    for c, text in zip(t.rows[0].cells, headers): c.text = text
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = text
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                prop = p._p.get_or_add_pPr(); bidi = OxmlElement('w:bidi'); prop.append(bidi)
                for r in p.runs:
                    r.font.name = 'DejaVu Sans'; r.font.size = Pt(9)
                    rtl = OxmlElement('w:rtl'); r._r.get_or_add_rPr().append(rtl)
    return t

# Current proposal is maintained in one Markdown source.
source = ROOT.parent / 'context' / 'proposal-working-draft.md'
for line in source.read_text().splitlines():
    if line.startswith('# '):
        para(line[2:], True, 22, True)
    elif line.startswith('## '):
        if line[3:] != 'דף שער':
            page()
        heading(line[3:])
    elif line.strip():
        para(line)

# Functional diagram: solid arrows represent data, power is separate.
im = Image.new('RGB', (1500, 1000), 'white')
draw = ImageDraw.Draw(im)
font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 25)
def box(x, y, w, h, lines):
    draw.rounded_rectangle((x,y,x+w,y+h), 15, fill='#e9f1fa', outline='#254465', width=3)
    for i, line in enumerate(lines):
        draw.text((x+w/2,y+18+i*34), line, font=font, fill='#16334d', anchor='mt', direction='rtl')
def arrow(a, b, both=False):
    import math
    draw.line((a,b), fill='#254465', width=3)
    def tip(start,end):
        angle=math.atan2(end[1]-start[1],end[0]-start[0])
        draw.polygon([end,(end[0]-15*math.cos(angle-.4),end[1]-15*math.sin(angle-.4)),(end[0]-15*math.cos(angle+.4),end[1]-15*math.sin(angle+.4))], fill='#254465')
    tip(a,b)
    if both: tip(b,a)
box(40,40,390,100,['טלפון: מפה, מצב ופקודות'])
box(540,40,420,100,['מחשב נייד: ניווט ומיפוי'])
box(1070,40,390,100,['שירות אינטרנט ו־AI','הערכה מייעצת בלבד'])
arrow((430,90),(540,90),True)
arrow((960,90),(1070,90),True)
box(180,350,470,115,['ESP: תנועה ומשוב','עצירה מקומית'])
box(850,350,470,115,['מצלמה עם מעבד עצמאי','חומרה ואישור טרם נקבעו'])
arrow((650,95),(415,350),True)
arrow((1085,350),(850,140))
box(30,650,370,100,['מקודדים וחיישני מרחק'])
box(455,650,310,100,['דרייברים ומנועים'])
box(950,650,370,100,['תמונות ואירועי חשד'])
arrow((215,650),(330,465))
arrow((510,465),(610,650))
arrow((1135,650),(1085,465))
box(420,850,660,100,['אספקת חשמל לרכיבי הרובוט','תקשורת ופרוטוקולים: לבחירה'])
im.save(ROOT/'robot-block-diagram.png')
page()
heading('תרשים זרימת מידע')
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(ROOT/'robot-block-diagram.png'), width=Cm(16.5))
para('התרשים פונקציונלי. כיווני המידע מוצגים בחצים; אספקת החשמל בנפרד. חיבור תפעולי בין שתי תתי־המערכות, הפרוטוקולים וחיבור בקר לאינטרנט טעונים תכנון ואימות לפי דרישות המנחה.')
doc.save(ROOT/'robot-project-proposal-draft.docx')
print(ROOT/'robot-project-proposal-draft.docx')
